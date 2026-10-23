"""Check a built deck: text density, small fonts, shapes off the slide and missing speaker notes; export slide previews with PowerPoint."""

import json
import re
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Pt

DENSITY = {"minimal": 45, "balanced": 80, "detailed": 120}
MIN_FONT = Pt(9)
LETTER_SPACING = re.compile(r"(?<=\S)   (?=\S)")


def lint(deck: Path, density: str, notes: bool) -> dict[str, object]:
    """Return the deck's problems per slide: words over the density limit, fonts under 9 pt, shapes past the edge, missing notes."""
    prs = Presentation(str(deck))
    width, height = Emu(prs.slide_width), Emu(prs.slide_height)
    problems, words = [], []
    for page, slide in enumerate(prs.slides, start=1):
        count = 0
        for shape in slide.shapes:
            if shape.left is not None and (shape.left < 0 or shape.top < 0 or shape.left + shape.width > width
                                           or shape.top + shape.height > height):
                problems.append(f"slide {page}: '{shape.name}' runs off the slide")
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                count += len(LETTER_SPACING.sub("", paragraph.text).split())
                sizes = [r.font.size for r in paragraph.runs if r.font.size is not None]
                if paragraph.text.strip() and any(s < MIN_FONT for s in sizes):
                    problems.append(f"slide {page}: text under 9 pt: {paragraph.text[:40]!r}")
        words.append(count)
        if count > DENSITY[density]:
            problems.append(f"slide {page}: {count} words, over {DENSITY[density]} for a {density} deck: cut text or turn it into a visual")
        if notes and not (slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip()):
            problems.append(f"slide {page}: no speaker notes")
    return {"ok": not problems, "slides": len(words), "words_per_slide": words, "problems": problems}


def preview(deck: Path, out: Path) -> list[str]:
    """Export every slide as a 1600x900 PNG with PowerPoint (Windows) and return the image paths."""
    if sys.platform != "win32":
        raise OSError("preview needs PowerPoint on Windows; open the deck and look at each slide instead")
    out.mkdir(parents=True, exist_ok=True)
    script = ("$app = New-Object -ComObject PowerPoint.Application; "
              f"$deck = $app.Presentations.Open('{deck.resolve()}', $true, $false, $false); "
              f"$deck.Export('{out.resolve()}', 'PNG', 1600, 900); $deck.Close(); $app.Quit()")
    subprocess.run(["powershell", "-NoProfile", "-Command", script], check=True, capture_output=True, timeout=300)
    images = sorted(out.glob("*.PNG"), key=lambda p: int("".join(c for c in p.stem if c.isdigit()) or 0))
    if not images:
        raise RuntimeError(f"PowerPoint exported no slides to {out}")
    return [p.as_posix() for p in images]


def main() -> None:
    """Usage: check_deck.py lint <deck.pptx> [minimal|balanced|detailed] [--notes] | preview <deck.pptx> <out folder>."""
    args = sys.argv[1:]
    if args[:1] == ["lint"] and len(args) >= 2 and set(args[2:]) <= {*DENSITY, "--notes"}:
        density = next((a for a in args[2:] if a in DENSITY), "balanced")
        result = lint(Path(args[1]), density, "--notes" in args)
        print(json.dumps(result, indent=2))
        sys.exit(0 if result["ok"] else 1)
    if args[:1] == ["preview"] and len(args) == 3:
        print(json.dumps(preview(Path(args[1]), Path(args[2])), indent=2))
        return
    raise SystemExit(main.__doc__)


if __name__ == "__main__":
    main()
