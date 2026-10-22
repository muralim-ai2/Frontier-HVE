"""NorthStar slide kit: a three-colour theme and python-pptx primitives for header bands, cards, pills, flows, layers, metrics and takeaway bars."""

from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.shapes.autoshape import Shape
from pptx.slide import Slide
from pptx.text.text import TextFrame
from pptx.util import Emu, Inches, Length, Pt

ACCENTS = {"green": "01A982", "blue": "0F6CBD", "teal": "00869B", "orange": "E8711C", "purple": "7A3E9D", "red": "C4314B"}
Line = tuple[str, float, bool, RGBColor, PP_ALIGN]


def tint(hex_code: str, share: float) -> RGBColor:
    """Return the colour mixed with white, keeping `share` of the colour."""
    channels = (int(hex_code[i:i + 2], 16) for i in (0, 2, 4))
    return RGBColor(*(round(c * share + 255 * (1 - share)) for c in channels))


@dataclass(frozen=True)
class Theme:
    """Three colours (ink for structure, one accent, slate for text), their tints, and the typeface."""

    ink: RGBColor
    accent: RGBColor
    slate: RGBColor
    ink_pale: RGBColor
    mist: RGBColor
    paper: RGBColor
    line: RGBColor
    accent_mist: RGBColor
    white: RGBColor
    font: str


def theme(accent: str = "green", ink: str = "1A2B34", slate: str = "62717A", font: str = "Segoe UI") -> Theme:
    """Return the NorthStar theme for an accent name from ACCENTS or a six-digit hex code."""
    code = ACCENTS.get(accent, accent).lstrip("#")
    if len(code) != 6:
        raise ValueError(f"accent {accent!r} is neither one of {sorted(ACCENTS)} nor a six-digit hex code")
    return Theme(ink=RGBColor.from_string(ink), accent=RGBColor.from_string(code), slate=RGBColor.from_string(slate),
                 ink_pale=tint(ink, 0.37), mist=tint(ink, 0.07), paper=tint(ink, 0.04), line=tint(ink, 0.16),
                 accent_mist=tint(code, 0.10), white=RGBColor(0xFF, 0xFF, 0xFF), font=font)


class Deck:
    """A 16:9 deck in the NorthStar grammar; footers with page numbers are added when it is saved."""

    def __init__(self, name: str, look: Theme) -> None:
        """Start an empty deck with the footer name and theme."""
        self.name, self.t = name, look
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.left, self.width = Inches(0.56), Emu(self.prs.slide_width - Inches(1.12))
        self.covers: set[int] = set()

    def ln(self, text: str, size: float, color: RGBColor | None = None, bold: bool = False, align: PP_ALIGN = PP_ALIGN.LEFT) -> Line:
        """Return one paragraph spec; colour defaults to slate."""
        return (text, size, bold, color or self.t.slate, align)

    def paint(self, shape: Shape, fill: RGBColor, edge: RGBColor | None = None) -> Shape:
        """Fill a shape flat (no shadow), with a 1 pt edge or none."""
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        if edge is None:
            shape.line.fill.background()
        else:
            shape.line.color.rgb = edge
            shape.line.width = Pt(1)
        shape.shadow.inherit = False
        return shape

    def write(self, frame: TextFrame, lines: list[Line], anchor: MSO_ANCHOR = MSO_ANCHOR.TOP) -> TextFrame:
        """Replace a text frame's text with the given paragraphs."""
        frame.word_wrap = True
        frame.clear()
        frame.vertical_anchor = anchor
        for i, (text, size, bold, color, align) in enumerate(lines):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            p.text, p.alignment, p.space_after = text, align, Pt(3)
            p.font.size, p.font.bold, p.font.name = Pt(size), bold, self.t.font
            p.font.color.rgb = color
        return frame

    def textbox(self, slide: Slide, left: Length, top: Length, width: Length, height: Length, lines: list[Line],
                anchor: MSO_ANCHOR = MSO_ANCHOR.TOP) -> Shape:
        """Add a borderless text box."""
        box = slide.shapes.add_textbox(left, top, width, height)
        self.write(box.text_frame, lines, anchor)
        return box

    def rect(self, slide: Slide, left: Length, top: Length, width: Length, height: Length, fill: RGBColor,
             edge: RGBColor | None = None, rounded: float | None = None) -> Shape:
        """Add a flat rectangle, rounded by the given corner share when set."""
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, left, top, width, height)
        if rounded:
            shape.adjustments[0] = rounded
        return self.paint(shape, fill, edge)

    def columns(self, count: int, gap: Length = Inches(0.20), left: Length | None = None, width: Length | None = None) -> tuple[list[Emu], Emu]:
        """Return the left edges and the width of `count` equal columns across the content width (or a given span)."""
        start, span = left if left is not None else self.left, width if width is not None else self.width
        col = Emu(int((span - gap * (count - 1)) / count))
        return [Emu(int(start + (col + gap) * i)) for i in range(count)], col

    def cover(self, title: str, subtitle: str, kicker: str) -> Slide:
        """Add an ink title slide: letter-spaced kicker, large title, subtitle and an accent rule."""
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.covers.add(slide.slide_id)
        self.rect(slide, 0, 0, self.prs.slide_width, self.prs.slide_height, self.t.ink)
        self.rect(slide, self.left, Inches(2.55), Inches(1.2), Inches(0.07), self.t.accent)
        self.textbox(slide, self.left, Inches(2.0), self.width, Inches(0.4), [self.ln(spaced(kicker), 11, self.t.accent, True)])
        self.textbox(slide, self.left, Inches(2.8), self.width, Inches(1.6), [self.ln(title, 40, self.t.white, True)])
        self.textbox(slide, self.left, Inches(4.5), self.width, Inches(1.0), [self.ln(subtitle, 16, self.t.ink_pale)])
        return slide

    def slide(self, eyebrow: str, title: str) -> Slide:
        """Add a content slide with the ink header band, accent rule, letter-spaced eyebrow and takeaway title."""
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.rect(slide, 0, 0, self.prs.slide_width, Inches(1.02), self.t.ink)
        self.rect(slide, 0, Inches(1.02), self.prs.slide_width, Inches(0.05), self.t.accent)
        self.textbox(slide, self.left, Inches(0.14), self.width, Inches(0.28), [self.ln(spaced(eyebrow), 9, self.t.ink_pale, True)])
        self.textbox(slide, self.left, Inches(0.42), self.width, Inches(0.5), [self.ln(title, 23, self.t.white, True)])
        return slide

    def label(self, slide: Slide, top: Length, text: str, left: Length | None = None, width: Length | None = None) -> Shape:
        """Add a letter-spaced section label; give it the full width of its column or it wraps."""
        return self.textbox(slide, left if left is not None else self.left, top, width if width is not None else self.width,
                            Inches(0.28), [self.ln(spaced(text), 9.5, self.t.slate, True)])

    def card(self, slide: Slide, left: Length, top: Length, width: Length, height: Length, lines: list[Line],
             fill: RGBColor | None = None, edge: RGBColor | None = None, rule: RGBColor | None = None,
             anchor: MSO_ANCHOR = MSO_ANCHOR.TOP) -> Shape:
        """Add a rounded card with an optional vertical rule on its left and its text inset."""
        box = self.rect(slide, left, top, width, height, fill or self.t.paper, edge or self.t.line, rounded=0.05)
        pad = Inches(0.22)
        if rule is not None:
            self.rect(slide, left + Inches(0.20), top + Inches(0.20), Inches(0.05), height - Inches(0.40), rule)
            pad = Inches(0.42)
        self.textbox(slide, left + pad, top + Inches(0.16), width - pad - Inches(0.20), height - Inches(0.30), lines, anchor)
        return box

    def highlight(self, slide: Slide, left: Length, top: Length, width: Length, height: Length, lines: list[Line]) -> Shape:
        """Add the slide's one emphasised card: accent tint, accent edge and accent rule, text centred vertically."""
        return self.card(slide, left, top, width, height, lines, self.t.accent_mist, self.t.accent, self.t.accent, MSO_ANCHOR.MIDDLE)

    def pill(self, slide: Slide, left: Length, top: Length, width: Length, height: Length, text: str, sub: str | None = None,
             strong: bool = False, size: float = 10.5) -> Shape:
        """Add a rounded chip with a bold head and an optional sub-line; `strong` gives it the accent."""
        box = self.rect(slide, left, top, width, height, self.t.accent_mist if strong else self.t.mist,
                        self.t.accent if strong else self.t.line, rounded=0.12)
        lines = [self.ln(text, size, self.t.accent if strong else self.t.ink, True, PP_ALIGN.CENTER)]
        lines += [self.ln(sub, size - 1.5, self.t.slate, False, PP_ALIGN.CENTER)] if sub else []
        self.write(box.text_frame, lines, MSO_ANCHOR.MIDDLE)
        box.text_frame.margin_left = box.text_frame.margin_right = Inches(0.06)
        return box

    def bar(self, slide: Slide, top: Length, text: str, height: Length = Inches(0.74), size: float = 14) -> Shape:
        """Add the slide's takeaway: a full-width ink bar with one bold white sentence."""
        box = self.rect(slide, self.left, top, self.width, height, self.t.ink, rounded=0.10)
        self.write(box.text_frame, [self.ln(text, size, self.t.white, True, PP_ALIGN.CENTER)], MSO_ANCHOR.MIDDLE)
        return box

    def flow(self, slide: Slide, top: Length, steps: list[tuple[str, str]], height: Length = Inches(1.24)) -> list[Shape]:
        """Add numbered process steps joined by accent chevrons; the last step is inked as the outcome."""
        lefts, width = self.columns(len(steps), Inches(0.24))
        boxes = []
        for i, ((head, sub), x) in enumerate(zip(steps, lefts)):
            last = i == len(steps) - 1
            box = self.rect(slide, x, top, width, height, self.t.ink if last else self.t.paper, None if last else self.t.line, rounded=0.06)
            self.write(box.text_frame, [self.ln(f"{i + 1:02d}", 9, self.t.accent, True, PP_ALIGN.CENTER),
                                        self.ln(head, 12.5, self.t.white if last else self.t.ink, True, PP_ALIGN.CENTER),
                                        self.ln(sub, 9, self.t.ink_pale if last else self.t.slate, False, PP_ALIGN.CENTER)],
                       MSO_ANCHOR.MIDDLE)
            box.text_frame.margin_left = box.text_frame.margin_right = Inches(0.06)
            if not last:
                chevron = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, x + width + Inches(0.01), top + Emu(int(height / 2)) - Inches(0.22),
                                                 Inches(0.22), Inches(0.44))
                self.paint(chevron, self.t.accent)
            boxes.append(box)
        return boxes

    def layers(self, slide: Slide, top: Length, rows: list[tuple[str, str]], height: Length = Inches(0.94),
               gap: Length = Inches(0.14), tag_width: Length = Inches(3.2)) -> None:
        """Add a stacked architecture: one band per layer with a coloured tag and its description; the top layer gets the accent."""
        for i, (name, description) in enumerate(rows):
            y = top + Emu(int((height + gap) * i))
            self.rect(slide, self.left, y, self.width, height, self.t.paper, self.t.line, rounded=0.06)
            tag = self.rect(slide, self.left + Inches(0.10), y + Inches(0.10), tag_width, height - Inches(0.20),
                            self.t.accent if i == 0 else self.t.ink, rounded=0.08)
            self.write(tag.text_frame, [self.ln(name, 12.5, self.t.white, True, PP_ALIGN.CENTER)], MSO_ANCHOR.MIDDLE)
            self.textbox(slide, self.left + tag_width + Inches(0.34), y + Inches(0.10), self.width - tag_width - Inches(0.56),
                         height - Inches(0.20), [self.ln(description, 10.5)], MSO_ANCHOR.MIDDLE)

    def metrics(self, slide: Slide, top: Length, items: list[tuple[str, str]], height: Length = Inches(1.3)) -> None:
        """Add metric tiles: a large accent number over a short caption."""
        lefts, width = self.columns(len(items))
        for (value, caption), x in zip(items, lefts):
            self.card(slide, x, top, width, height, [self.ln(value, 30, self.t.accent, True, PP_ALIGN.CENTER),
                                                     self.ln(caption, 10.5, self.t.slate, False, PP_ALIGN.CENTER)],
                      fill=self.t.white, anchor=MSO_ANCHOR.MIDDLE)

    def notes(self, slide: Slide, text: str) -> None:
        """Set the slide's speaker notes."""
        slide.notes_slide.notes_text_frame.text = text

    def save(self, path: Path) -> Path:
        """Add the footer (deck name, page / total) to every slide except covers and save the deck."""
        total = len(self.prs.slides)
        for page, slide in enumerate(self.prs.slides, start=1):
            if slide.slide_id not in self.covers:
                self.textbox(slide, self.left, self.prs.slide_height - Inches(0.36), self.width, Inches(0.28),
                             [self.ln(f"{self.name}          {page} / {total}", 8.5)])
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(path))
        return path


def spaced(text: str) -> str:
    """Return text in letter-spaced capitals for eyebrows and labels."""
    return "   ".join(text.upper())
